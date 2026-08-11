"""Validate and render the official weekly-minutes DOCX template."""

from __future__ import annotations

from datetime import datetime
import logging
from pathlib import Path
import shutil
import tempfile
from typing import Any

from docx import Document
from docxtpl import DocxTemplate

from .docx_merge import compose_person_docx, insert_docx_body_before
from .people import PeopleDirectory, PeopleStore, ensure_store
from .repository import SubmissionContent


LOGGER = logging.getLogger(__name__)
RICH_MARKER_PREFIX = "__MEETING_RICH__"
RICH_MARKER_SUFFIX = "__"


class MinutesTemplateError(ValueError):
    """Raised when the configured Word template cannot safely be rendered."""


def numbered_contents(contents: tuple[str, ...]) -> str:
    if not contents:
        return "本周未提交"
    return "\n".join(f"{index}. {content}" for index, content in enumerate(contents, 1))


def _rich_marker(template_key: str) -> str:
    return f"{RICH_MARKER_PREFIX}{template_key}{RICH_MARKER_SUFFIX}"


class MinutesDocumentRenderer:
    def __init__(
        self,
        *,
        template_path: str | Path,
        output_dir: str | Path,
        people: PeopleDirectory | PeopleStore,
        data_dir: str | Path | None = None,
    ) -> None:
        self.template_path = Path(template_path).resolve()
        self.output_dir = Path(output_dir).resolve()
        self._people_store = ensure_store(people)
        self.data_dir = (
            Path(data_dir).resolve()
            if data_dir is not None
            else self.output_dir.parent.resolve()
        )
        self.validate_template()

    @property
    def people(self) -> PeopleDirectory:
        return self._people_store.directory

    def validate_template(
        self, directory: PeopleDirectory | None = None
    ) -> frozenset[str]:
        """Validate placeholders against ``directory`` (default: active people)."""

        people = directory if directory is not None else self.people
        if not self.template_path.is_file():
            raise MinutesTemplateError(f"Word 模板不存在：{self.template_path}")
        try:
            document = DocxTemplate(self.template_path)
            variables = frozenset(document.get_undeclared_template_variables())
        except Exception as exc:
            raise MinutesTemplateError(f"Word 模板无法打开：{exc}") from exc

        required = {person.template_key for person in people.enabled_people}
        missing = sorted(required - variables)
        if missing:
            raise MinutesTemplateError(
                "Word 模板缺少启用人员占位符：" + "、".join(missing)
            )
        return variables

    def output_path(self, period: str, version: int, generated_at: datetime) -> Path:
        year, week = period.split("-W", 1)
        timestamp = generated_at.strftime("%Y%m%d_%H%M%S")
        filename = f"{year}年第{int(week)}周周例会纪要_v{version}_{timestamp}.docx"
        return self.output_dir / filename

    def _resolve_source(self, relative: str | None) -> Path | None:
        if not relative:
            return None
        path = Path(relative)
        if not path.is_absolute():
            path = self.data_dir / path
        return path if path.is_file() else None

    def _normalize_items(
        self,
        raw_items: tuple[Any, ...],
    ) -> tuple[SubmissionContent, ...]:
        if not raw_items:
            return ()
        if isinstance(raw_items[0], SubmissionContent):
            return raw_items  # type: ignore[return-value]
        return tuple(
            SubmissionContent(parsed_content=str(text)) for text in raw_items
        )

    def _prepare_person_value(
        self,
        *,
        items: tuple[SubmissionContent, ...],
        work_dir: Path,
        template_key: str,
        rich_files: dict[str, Path],
    ) -> str:
        if not items:
            return "本周未提交"

        needs_rich = any(
            item.source_relative_path and item.message_type == "docx"
            for item in items
        )
        if not needs_rich:
            return numbered_contents(tuple(item.parsed_content for item in items))

        composed_items: list[tuple[str, Path | None]] = []
        for item in items:
            source = (
                self._resolve_source(item.source_relative_path)
                if item.message_type == "docx"
                else None
            )
            composed_items.append((item.parsed_content, source))

        composed_path = work_dir / f"{template_key}.docx"
        try:
            compose_person_docx(items=composed_items, output_path=composed_path)
            marker = _rich_marker(template_key)
            rich_files[marker] = composed_path
            return marker
        except Exception:
            LOGGER.exception(
                "人员 %s 的富文本合成失败，回退为文字摘要", template_key
            )
            return numbered_contents(tuple(item.parsed_content for item in items))

    def _inject_rich_sections(
        self, document_path: Path, rich_files: dict[str, Path]
    ) -> None:
        """Replace marker paragraphs with composed DOCX body content."""

        if not rich_files:
            return
        document = Document(str(document_path))
        for paragraph in list(document.paragraphs):
            marker = paragraph.text.strip()
            source = rich_files.get(marker)
            if source is None:
                continue
            insert_at = paragraph._element
            parent = insert_at.getparent()
            if parent is None:
                continue
            try:
                insert_docx_body_before(document, source, insert_at)
            except Exception:
                LOGGER.exception("注入富文本失败，保留文字标记：%s", marker)
                continue
            parent.remove(insert_at)
        document.save(document_path)

    def render(
        self,
        *,
        period: str,
        version: int,
        generated_at: datetime,
        contents: dict[str, tuple[str, ...]]
        | dict[str, tuple[SubmissionContent, ...]],
    ) -> Path:
        self.output_dir.mkdir(parents=True, exist_ok=True)
        output_path = self.output_path(period, version, generated_at)
        temporary = output_path.with_name(f".{output_path.name}.tmp.docx")
        work_dir = Path(tempfile.mkdtemp(prefix="meeting-compose-"))
        rich_files: dict[str, Path] = {}
        try:
            context: dict[str, str] = {}
            for person in self.people.people:
                items = self._normalize_items(
                    contents.get(person.template_key, ())  # type: ignore[arg-type]
                )
                context[person.template_key] = self._prepare_person_value(
                    items=items,
                    work_dir=work_dir,
                    template_key=person.template_key,
                    rich_files=rich_files,
                )
            document = DocxTemplate(self.template_path)
            document.render(context, autoescape=True)
            document.save(temporary)
            self._inject_rich_sections(temporary, rich_files)
            temporary.replace(output_path)
        except Exception as exc:
            temporary.unlink(missing_ok=True)
            raise MinutesTemplateError(f"Word 纪要生成失败：{exc}") from exc
        finally:
            shutil.rmtree(work_dir, ignore_errors=True)
        return output_path
